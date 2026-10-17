"""
src/export_standard_deliverable.py
Generates the exact 6-column Deliverable 1 spreadsheet and CSV:
  Columns: NPI, TIN, Address, Billing Code, Specialty, Rate
As specified in official HiLabs Hackathon Deliverable 1 requirements.
"""
from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.utils import load_config, setup_logging

logger = logging.getLogger(__name__)


def export_standard_deliverables(
    benchmark_csv_path: str = "outputs/c6_curated_benchmark_rates.csv",
    output_xlsx_path: str = "outputs/standard_curated_rates.xlsx",
    output_csv_path: str = "outputs/c6_curated_standard_rates.csv",
    submission_xlsx_path: str = "submission/standard_curated_rates.xlsx",
    submission_csv_path: str = "submission/c6_curated_standard_rates.csv",
) -> int:
    """Generate the exact 6-column Deliverable 1 files."""
    records: List[List[object]] = []
    with open(benchmark_csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r in reader:
            parts = [
                str(r.get("address_line_1", "") or "").strip(),
                str(r.get("address_line_2", "") or "").strip(),
            ]
            street = ", ".join([p for p in parts if p])
            city_st_zip = f"{str(r.get('city', '') or '').strip()}, {str(r.get('state', '') or '').strip()} {str(r.get('postal_code', '') or '').strip()}".strip()
            full_addr = f"{street}, {city_st_zip}".strip(", ") if street else city_st_zip
            records.append([
                str(r["npi"]),
                str(r["tin_value"]),
                full_addr,
                str(r["billing_code"]),
                str(r.get("specialty_description", "") or ""),
                float(r["negotiated_rate"]),
            ])

    # 1. Write CSV files
    for p in (output_csv_path, submission_csv_path):
        Path(p).parent.mkdir(parents=True, exist_ok=True)
        with open(p, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["NPI", "TIN", "Address", "Billing Code", "Specialty", "Rate"])
            w.writerows(records)
        logger.info("Saved standard 6-column CSV: %s", p)

    # 2. Write Excel files
    fill_hdr = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    font_hdr = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fill_zebra = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")
    font_cell = Font(name="Segoe UI", size=10)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for p in (output_xlsx_path, submission_xlsx_path):
        Path(p).parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Standard Rate File"
        ws.views.sheetView[0].showGridLines = True
        headers = ["NPI", "TIN", "Address", "Billing Code", "Specialty", "Rate"]
        ws.append(headers)

        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=c_idx)
            c.fill = fill_hdr
            c.font = font_hdr
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28

        for r_idx, r in enumerate(records, start=2):
            ws.append(r)
            is_even = (r_idx % 2 == 0)
            for c_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = font_cell
                cell.border = thin_border
                if is_even:
                    cell.fill = fill_zebra
                if c_idx in (1, 2, 4):
                    cell.number_format = "@"
                    cell.alignment = Alignment(horizontal="center")
                elif c_idx == 6:
                    cell.number_format = "$#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                    cell.font = font_bold
                else:
                    cell.alignment = Alignment(horizontal="left")

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:F{len(records) + 1}"
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or "")) for cell in col[:100])
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

        wb.save(p)
        logger.info("Saved standard 6-column Excel: %s", p)

    return len(records)


def main() -> None:
    setup_logging(log_level="INFO")
    cfg = load_config()
    export_standard_deliverables(
        benchmark_csv_path=cfg["c6_curated_benchmark_rates_csv"],
        output_xlsx_path="outputs/standard_curated_rates.xlsx",
        output_csv_path="outputs/c6_curated_standard_rates.csv",
        submission_xlsx_path="submission/standard_curated_rates.xlsx",
        submission_csv_path="submission/c6_curated_standard_rates.csv",
    )


if __name__ == "__main__":
    main()
