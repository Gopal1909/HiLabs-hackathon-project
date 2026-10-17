"""
tests/test_c6_rate_curator.py
Comprehensive test suite for Phase 6 (C6) Rate Curation and Excel Deliverable.
"""
from __future__ import annotations

import csv
import json
import os
from pathlib import Path
from typing import Any, Dict

import openpyxl
import pytest

from src.c6_rate_curator import (
    build_curated_record,
    curate_and_export_rates,
    load_benchmark_tables,
    load_c5_provider_crosswalk,
)
from src.utils import load_config


@pytest.fixture(scope="module")
def cfg() -> Dict[str, Any]:
    return load_config()


def test_load_c5_provider_crosswalk(cfg: Dict[str, Any]):
    """Verify C5 crosswalk loads with expected key tuple and attributes."""
    c5_path = cfg["c5_provider_address_crosswalk_csv"]
    c5_map = load_c5_provider_crosswalk(c5_path)
    assert len(c5_map) == 5672
    sample_key = next(iter(c5_map))
    assert len(sample_key) == 3
    assert all(isinstance(k, str) for k in sample_key)
    rec = c5_map[sample_key]
    assert "address_line_1" in rec
    assert "verification_status" in rec
    assert "confidence_score" in rec
    assert "primary_taxonomy" in rec


def test_load_benchmark_tables(cfg: Dict[str, Any]):
    """Verify loading of benchmark weights tables."""
    bench_path = cfg["benchmark_weights_xlsx"]
    tax_to_spec, spec_weights, spec_proc_weights, proc_max_weights = load_benchmark_tables(bench_path)

    assert len(tax_to_spec) > 200
    assert len(spec_weights) > 40
    assert len(spec_proc_weights) > 900
    assert len(proc_max_weights) > 500

    # Check known sample codes
    assert "207K00000X" in tax_to_spec
    assert tax_to_spec["207K00000X"][0] == "007"
    assert "PCP" in spec_weights
    assert spec_weights["PCP"][1] > 0
    assert ("007", "J1559") in spec_proc_weights
    assert "99214" in proc_max_weights


def test_build_curated_record_exact_weight():
    """Verify record enrichment with exact (spec_cd, proc_cd) weight."""
    c4_row = {
        "npi": "1234567890",
        "tin_type": "ein",
        "tin_value": "12-3456789",
        "provider_business_name": "Test Clinic",
        "provider_group_id": "1",
        "billing_code": "99214",
        "billing_code_type": "CPT",
        "billing_code_description": "Office Visit",
        "billing_code_modifier": "",
        "billing_class": "professional",
        "setting": "outpatient",
        "service_codes": "11",
        "negotiated_rate": "150.00",
        "rate_selection_flag": "UNIQUE_RATE",
        "alternative_rate_values": "",
        "ppo_applicability": "CONFIRMED_PPO",
        "is_benchmark_code": "Y",
        "source_file": "test.json",
    }
    prov = {
        "npi_type": "Type 1 (Individual)",
        "provider_name": "DR. SMITH",
        "ptd_business_name": "Test Clinic",
        "address_line_1": "100 MAIN ST",
        "address_line_2": "STE 200",
        "city": "SACRAMENTO",
        "state": "CA",
        "postal_code": "95814",
        "address_purpose": "LOCATION",
        "verification_status": "VERIFIED_CURRENT",
        "confidence_score": "0.95",
        "evidence_sources": "NPPES;ROSTER",
        "is_stale_nppes_risk": "N",
        "primary_taxonomy": "207Q00000X",  # Family Medicine -> PCP
        "specialty_description": "Family Medicine",
    }
    tax_to_spec = {"207Q00000X": ("PCP", "Primary Care")}
    spec_weights = {"PCP": ("Primary Care", 3932892)}
    spec_proc_weights = {("PCP", "99214"): (1, 692869553)}
    proc_max_weights = {"99214": (1, 692869553)}

    rec = build_curated_record(
        c4_row, prov, tax_to_spec, spec_weights, spec_proc_weights, proc_max_weights
    )

    assert rec["npi"] == "1234567890"
    assert rec["specialty_code"] == "PCP"
    assert rec["specialty_description"] == "Primary Care"
    assert rec["benchmark_rank"] == 1
    assert rec["benchmark_procedure_weight"] == 692869553
    assert rec["benchmark_specialty_weight"] == 3932892
    assert rec["weight_match_level"] == "SPECIALTY_AND_PROCEDURE"
    assert rec["negotiated_rate"] == 150.00
    assert rec["address_verification_status"] == "VERIFIED_CURRENT"


def test_build_curated_record_procedure_fallback():
    """Verify fallback procedure-level weight when specialty has no direct pair."""
    c4_row = {
        "npi": "1234567890",
        "tin_type": "ein",
        "tin_value": "12-3456789",
        "provider_business_name": "General Acute Hospital",
        "provider_group_id": "1",
        "billing_code": "99214",
        "billing_code_type": "CPT",
        "billing_code_description": "Office Visit",
        "billing_code_modifier": "",
        "billing_class": "institutional",
        "setting": "outpatient",
        "service_codes": "22",
        "negotiated_rate": "220.00",
        "rate_selection_flag": "UNIQUE_RATE",
        "alternative_rate_values": "",
        "ppo_applicability": "PROBABLE_PPO",
        "is_benchmark_code": "Y",
        "source_file": "test.json",
    }
    prov = {
        "npi_type": "Type 2 (Organization)",
        "provider_name": "COMMUNITY HOSPITAL",
        "ptd_business_name": "Community Hospital Clinic",
        "address_line_1": "500 HOSPITAL WAY",
        "address_line_2": "",
        "city": "FRESNO",
        "state": "CA",
        "postal_code": "93720",
        "address_purpose": "LOCATION",
        "verification_status": "VERIFIED_CURRENT",
        "confidence_score": "0.95",
        "evidence_sources": "NPPES;FACILITY",
        "is_stale_nppes_risk": "N",
        "primary_taxonomy": "282N00000X",  # General Acute Care Hospital (no direct SPEC_CD)
        "specialty_description": "General Acute Care Hospital",
    }
    tax_to_spec = {}
    spec_weights = {}
    spec_proc_weights = {}
    proc_max_weights = {"99214": (1, 692869553)}

    rec = build_curated_record(
        c4_row, prov, tax_to_spec, spec_weights, spec_proc_weights, proc_max_weights
    )

    assert rec["specialty_code"] == ""
    assert rec["specialty_description"] == "General Acute Care Hospital"
    assert rec["benchmark_rank"] == 1
    assert rec["benchmark_procedure_weight"] == 692869553
    assert rec["benchmark_specialty_weight"] == ""
    assert rec["weight_match_level"] == "PROCEDURE_LEVEL_FALLBACK"


def test_build_curated_record_non_benchmark():
    """Verify non-benchmark codes receive no weights."""
    c4_row = {
        "npi": "1234567890",
        "tin_type": "ein",
        "tin_value": "12-3456789",
        "provider_business_name": "Lab",
        "provider_group_id": "1",
        "billing_code": "0202U",
        "billing_code_type": "HCPCS",
        "billing_code_description": "SARS-COV-2 Assay",
        "billing_code_modifier": "",
        "billing_class": "professional",
        "setting": "outpatient",
        "service_codes": "81",
        "negotiated_rate": "1463.00",
        "rate_selection_flag": "UNIQUE_RATE",
        "alternative_rate_values": "",
        "ppo_applicability": "PROBABLE_PPO",
        "is_benchmark_code": "N",
        "source_file": "test.json",
    }
    prov = {
        "npi_type": "Type 2 (Organization)",
        "provider_name": "LAB CORP",
        "ptd_business_name": "LabCorp Elk Grove",
        "address_line_1": "1000 LAB WAY",
        "address_line_2": "",
        "city": "ELK GROVE",
        "state": "CA",
        "postal_code": "95758",
        "address_purpose": "LOCATION",
        "verification_status": "VERIFIED_CURRENT",
        "confidence_score": "0.95",
        "evidence_sources": "NPPES",
        "is_stale_nppes_risk": "N",
        "primary_taxonomy": "291U00000X",
        "specialty_description": "Clinical Medical Laboratory",
    }
    rec = build_curated_record(c4_row, prov, {}, {}, {}, {})
    assert rec["is_benchmark_code"] == "N"
    assert rec["benchmark_rank"] == ""
    assert rec["benchmark_procedure_weight"] == ""
    assert rec["weight_match_level"] == "NONE"


def test_curated_benchmark_csv_deliverable(cfg: Dict[str, Any]):
    """Verify outputs/c6_curated_benchmark_rates.csv exists and has exactly 2,418 rows."""
    bench_csv = cfg["c6_curated_benchmark_rates_csv"]
    assert os.path.isfile(bench_csv), f"Missing {bench_csv}"

    with open(bench_csv, "r", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        rows = list(reader)
        assert len(rows) == 2418
        assert "npi" in reader.fieldnames
        assert "tin_value" in reader.fieldnames
        assert "address_line_1" in reader.fieldnames
        assert "city" in reader.fieldnames
        assert "state" in reader.fieldnames
        assert "postal_code" in reader.fieldnames
        assert "negotiated_rate" in reader.fieldnames
        assert "benchmark_rank" in reader.fieldnames
        assert "benchmark_procedure_weight" in reader.fieldnames

        # Verify rate values are strictly positive floats
        for r in rows:
            rate = float(r["negotiated_rate"])
            assert rate > 0.0
            assert r["ppo_applicability"] in ("CONFIRMED_PPO", "PROBABLE_PPO")


def test_curated_ppo_csv_deliverable(cfg: Dict[str, Any]):
    """Verify outputs/c6_curated_ppo_rates.csv exists, is non-empty, and has expected schema."""
    full_csv = cfg["c6_curated_ppo_rates_csv"]
    assert os.path.isfile(full_csv), f"Missing {full_csv}"
    assert os.path.getsize(full_csv) > 100_000_000  # Should be ~400 MB

    with open(full_csv, "r", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        assert len(reader.fieldnames) == 37
        first_row = next(reader)
        assert first_row["npi"] != ""
        assert first_row["tin_value"] != ""
        assert first_row["ppo_applicability"] in ("CONFIRMED_PPO", "PROBABLE_PPO")


def test_excel_deliverable_structure(cfg: Dict[str, Any]):
    """Verify curated_ppo_rates.xlsx has all 4 expected tabs, headers, and correct counts."""
    excel_path = cfg["c6_curated_ppo_rates_xlsx"]
    assert os.path.isfile(excel_path), f"Missing {excel_path}"

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    expected_sheets = ["Executive Summary", "Benchmark Rates", "PPO Rate Cross-Section", "Data Dictionary"]
    assert wb.sheetnames == expected_sheets

    # Tab 1: Executive Summary
    ws_sum = wb["Executive Summary"]
    assert ws_sum["B2"].value == "HiLabs Hackathon — Data Pipeline Curation for Provider Rates"
    assert ws_sum.views.sheetView[0].showGridLines is True

    # Tab 2: Benchmark Rates
    ws_bench = wb["Benchmark Rates"]
    assert ws_bench.max_row == 2419  # Header + 2,418 rows
    assert ws_bench.max_column == 26
    assert ws_bench.cell(row=1, column=1).value == "NPI"
    assert ws_bench.cell(row=1, column=21).value == "Negotiated Rate ($)"

    # Tab 3: PPO Rate Cross-Section
    ws_samp = wb["PPO Rate Cross-Section"]
    assert ws_samp.max_row == 5001  # Header + 5,000 rows
    assert ws_samp.max_column == 26

    # Tab 4: Data Dictionary
    ws_dict = wb["Data Dictionary"]
    assert ws_dict.max_row >= 30
    assert ws_dict.cell(row=5, column=2).value == "Field Name"


def test_curation_summary_json(cfg: Dict[str, Any]):
    """Verify outputs/c6_curation_summary.json structure and mathematical consistency."""
    summary_path = cfg["c6_curation_summary_json"]
    assert os.path.isfile(summary_path)

    with open(summary_path, "r", encoding="utf-8") as fh:
        summary = json.load(fh)

    vol = summary["curation_volume"]
    assert vol["total_extracted_c4_rows"] == 1511720
    assert vol["curated_ppo_rates"] == 857206
    assert vol["curated_benchmark_rates"] == 2418
    assert vol["excluded_non_ppo_rows"] == 654514
    assert vol["curated_ppo_rates"] + vol["excluded_non_ppo_rows"] == vol["total_extracted_c4_rows"]

    # Verify address verification rate is > 99%
    addr = summary["address_verification_breakdown"]
    assert addr["verified_and_corroborated_pct"] > 99.0

    # Rate statistics sanity
    bench_rates = summary["rate_distributions"]["benchmark_procedure_rates"]
    assert bench_rates["count"] == 2418
    assert bench_rates["min"] > 0
    assert bench_rates["min"] <= bench_rates["median"] <= bench_rates["max"]
    assert bench_rates["mean"] > 0
